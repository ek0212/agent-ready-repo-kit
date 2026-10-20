# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0"]
# ///
"""Deck helpers for the presentation-builder skill.

Subcommands:
  validate DRAFT.md               Check draft flow, 5 candidate visuals, and spoken notes per slide.
  selections DRAFT.md             Read user-pruned draft into selections JSON.
  inspect TEMPLATE.pptx           Print template fonts, sizes, colors, layouts, and icons.
  check FINAL.pptx TEMPLATE.pptx  Verify final deck uses only template styling and assets.
"""

import argparse
import hashlib
import json
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

import constants as C

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
BULLET_PATTERN = re.compile(r"^[-*]\s+")
BOLD_PATTERN = re.compile(r"\*\*[^*]+\*\*")
SECTION_PATTERN = re.compile(r"\(([^)]+)\)")
WORD_PATTERN = re.compile(r"[a-z']+")


def count_words(text):
    return len(text.split())


def split_sentences(text):
    return [s for s in re.split(r"(?<=[.!?])\s+", text.strip()) if s]


def count_sentences(text):
    return len(split_sentences(text))


def spoken_text(point):
    return point.replace(C.BOLD_MARKER, "").replace(C.PAUSE_MARKER, "").strip()


def start_slide(line):
    label, _, heading = line.partition(":")
    section = SECTION_PATTERN.search(label)
    return {"heading": heading.strip(), "section": section.group(1).strip() if section else "", "visuals": [], "notes": []}


def section_for(line):
    lowered = line.lower()
    if lowered == C.MD_VISUALS_HEADING.lower():
        return "visuals"
    if lowered == C.MD_NOTES_HEADING.lower():
        return "notes"
    return None


def add_line(slide, section, raw_line):
    line = raw_line.strip()
    if section == "visuals" and BULLET_PATTERN.match(line):
        slide["visuals"].append(BULLET_PATTERN.sub("", line).strip())
    elif section == "notes":
        slide["notes"].append(raw_line)


def parse_draft(text):
    slides, section = [], None
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if line.startswith(C.MD_SLIDE_PREFIX):
            slides.append(start_slide(line))
            section = None
        elif slides and section_for(line):
            section = section_for(line)
        elif slides:
            add_line(slides[-1], section, raw_line)
    for slide in slides:
        slide["notes"] = "\n".join(slide["notes"]).strip()
    return slides


def note_lines(notes):
    return [line.strip() for line in notes.splitlines() if line.strip()]


def strip_bullet(line):
    return BULLET_PATTERN.sub("", line).strip()


def validate_draft(slides):
    errors, warnings = [], []
    for n, slide in enumerate(slides, start=1):
        errors += validate_slide_text(n, slide)
        warnings += notes_warnings(n, note_points(slide["notes"]))
    flow_errors, flow_warnings = validate_flow([slide["section"] for slide in slides])
    errors, warnings = errors + flow_errors, warnings + flow_warnings
    if slides and not any(point.startswith(C.TAKEAWAY_PREFIX) for point in note_points(slides[-1]["notes"])):
        errors.append(C.MSG_NO_TAKEAWAY)
    return errors, warnings


def note_points(notes):
    return [strip_bullet(line) for line in note_lines(notes)]


def notes_errors(n, notes):
    if not notes:
        return [C.MSG_NOTES_EMPTY.format(n=n)]
    if not all(BULLET_PATTERN.match(line) for line in note_lines(notes)):
        return [C.MSG_NOTES_NOT_BULLETS.format(n=n)]
    errors = []
    for p, point in enumerate(note_points(notes), start=1):
        sentences = count_sentences(spoken_text(point))
        if sentences > C.MAX_SENTENCES_PER_NOTE_POINT:
            errors.append(C.MSG_POINT_SENTENCES.format(n=n, p=p, count=sentences, limit=C.MAX_SENTENCES_PER_NOTE_POINT))
    return errors


def validate_slide_text(n, slide):
    errors = []
    heading_words = count_words(slide["heading"])
    if heading_words > C.MAX_HEADING_WORDS:
        errors.append(C.MSG_HEADING_WORDS.format(n=n, count=heading_words, limit=C.MAX_HEADING_WORDS))
    if len(slide["visuals"]) != C.VISUALS_PER_DRAFT_SLIDE:
        errors.append(C.MSG_VISUAL_COUNT.format(n=n, expected=C.VISUALS_PER_DRAFT_SLIDE, actual=len(slide["visuals"])))
    for v, visual in enumerate(slide["visuals"], start=1):
        sentences = count_sentences(visual)
        if sentences > C.MAX_SENTENCES_PER_VISUAL:
            errors.append(C.MSG_VISUAL_SENTENCES.format(n=n, v=v, count=sentences, limit=C.MAX_SENTENCES_PER_VISUAL))
    return errors + notes_errors(n, slide["notes"])


def notes_warnings(n, points):
    if not points:
        return []
    warnings = []
    if not C.NOTE_POINTS_MIN <= len(points) <= C.NOTE_POINTS_MAX:
        warnings.append(C.MSG_POINT_COUNT.format(n=n, count=len(points), low=C.NOTE_POINTS_MIN, high=C.NOTE_POINTS_MAX))
    words = sum(count_words(point.replace(C.BOLD_MARKER, "")) for point in points)
    if not C.NOTES_WORDS_MIN <= words <= C.NOTES_WORDS_MAX:
        warnings.append(C.MSG_NOTES_LENGTH.format(n=n, count=words, low=C.NOTES_WORDS_MIN, high=C.NOTES_WORDS_MAX))
    if not any(BOLD_PATTERN.search(point) for point in points):
        warnings.append(C.MSG_NO_BOLD.format(n=n))
    return warnings + spoken_voice_warnings(n, points)


def spoken_voice_warnings(n, points):
    warnings = []
    for p, point in enumerate(points, start=1):
        for sentence in split_sentences(spoken_text(point)):
            if count_words(sentence) > C.MAX_SPOKEN_SENTENCE_WORDS:
                warnings.append(C.MSG_LONG_SENTENCE.format(n=n, p=p, count=count_words(sentence), limit=C.MAX_SPOKEN_SENTENCE_WORDS))
    text = " ".join(spoken_text(point) for point in points).lower()
    padded = f" {' '.join(WORD_PATTERN.findall(text))} "
    warnings += [C.MSG_TALKING_DOWN.format(n=n, value=phrase) for phrase in C.TALKING_DOWN_PHRASES if f" {phrase} " in padded]
    if not C.AUDIENCE_WORDS & set(WORD_PATTERN.findall(text.replace("’", "'"))):
        warnings.append(C.MSG_NO_AUDIENCE.format(n=n))
    return warnings


def validate_flow(sections):
    if not sections:
        return [], []
    allowed = " -> ".join(C.SECTIONS)
    errors = [C.MSG_SECTION_UNKNOWN.format(n=n, value=s, allowed=allowed) for n, s in enumerate(sections, start=1) if s not in C.SECTIONS]
    if errors:
        return errors, []
    order = [C.SECTIONS.index(s) for s in sections]
    errors += [C.MSG_SECTION_ORDER.format(n=n, value=sections[n - 1], allowed=allowed) for n in range(2, len(order) + 1) if order[n - 1] < order[n - 2]]
    if sections[0] != C.SECTIONS[0]:
        errors.append(C.MSG_SECTION_FIRST.format(value=C.SECTIONS[0]))
    if sections[-1] != C.SECTIONS[-1]:
        errors.append(C.MSG_SECTION_LAST.format(value=C.SECTIONS[-1]))
    warnings = [C.MSG_SECTION_MISSING.format(value=s) for s in C.SECTIONS if s not in sections]
    return errors, warnings


def read_selections(slides):
    issues = [C.MSG_NO_SELECTION.format(n=n) for n, slide in enumerate(slides, start=1) if not slide["visuals"]]
    selected = [{"heading": s["heading"], "selected_visuals": s["visuals"], "notes": s["notes"]} for s in slides]
    return {"slides": selected, "issues": issues}


def load_draft(path):
    with open(path, encoding="utf-8") as handle:
        return parse_draft(handle.read())


def theme_xml(prs):
    return etree.fromstring(prs.slide_masters[0].part.part_related_by(RT.THEME).blob)


def theme_colors(theme):
    colors = {}
    for slot in theme.find(".//a:clrScheme", NS):
        value = slot[0].get("lastClr") or slot[0].get("val")
        colors[etree.QName(slot).localname] = value.upper()
    return colors


def theme_fonts(theme):
    return {kind: theme.find(f".//a:fontScheme/a:{kind}/a:latin", NS).get("typeface") for kind in ("majorFont", "minorFont")}


def scan_styles(xml):
    fonts = {el.get("typeface") for el in xml.iterfind(".//a:latin", NS)} - C.THEME_FONT_REFS
    sizes = {int(el.get("sz")) / C.SZ_UNITS_PER_PT for el in xml.iter() if el.get("sz") and etree.QName(el).localname in ("rPr", "defRPr", "endParaRPr")}
    colors = {el.get("val").upper() for el in xml.iterfind(".//a:srgbClr", NS)}
    return fonts, sizes, colors


def image_hashes(part):
    return {hashlib.sha1(rel.target_part.blob).hexdigest(): rel.target_part.partname for rel in part.rels.values() if rel.reltype == RT.IMAGE}


def template_parts(prs):
    for master in prs.slide_masters:
        yield "master", master
        for layout in master.slide_layouts:
            yield f"layout '{layout.name}'", layout
    for n, slide in enumerate(prs.slides, start=1):
        yield f"slide {n}", slide


ICON_KINDS = {MSO_SHAPE_TYPE.PICTURE: "picture", MSO_SHAPE_TYPE.GROUP: "group", MSO_SHAPE_TYPE.FREEFORM: "freeform"}


def safe_shape_type(shape):
    try:
        return shape.shape_type
    except NotImplementedError:
        return None


def icon_entries(source, container):
    typed = [(s, safe_shape_type(s)) for s in container.shapes]
    return [{"source": source, "name": s.name, "kind": ICON_KINDS[t], "shape_id": s.shape_id} for s, t in typed if t in ICON_KINDS]


def layout_entries(prs):
    return [
        {"name": layout.name, "placeholders": [{"idx": ph.placeholder_format.idx, "type": str(ph.placeholder_format.type), "name": ph.name} for ph in layout.placeholders]}
        for layout in prs.slide_layouts
    ]


def inspect_template(path):
    prs = Presentation(path)
    theme = theme_xml(prs)
    fonts, sizes, colors, images, icons = set(theme_fonts(theme).values()), set(), set(theme_colors(theme).values()), {}, []
    for source, container in template_parts(prs):
        part_fonts, part_sizes, part_colors = scan_styles(container.part._element)
        fonts |= part_fonts
        sizes |= part_sizes
        colors |= part_colors
        images |= image_hashes(container.part)
        icons += icon_entries(source, container)
    return {
        "theme_fonts": theme_fonts(theme),
        "theme_colors": theme_colors(theme),
        "fonts": sorted(fonts),
        "sizes_pt": sorted(sizes),
        "colors": sorted(colors),
        "layouts": layout_entries(prs),
        "icons": icons,
        "images": {h: str(p) for h, p in images.items()},
    }


def check_slide(n, slide, allowed):
    fonts, sizes, colors = scan_styles(slide.part._element)
    issues = [C.MSG_FONT.format(n=n, value=f) for f in sorted(fonts - set(allowed["fonts"]))]
    issues += [C.MSG_SIZE.format(n=n, value=s) for s in sorted(sizes - set(allowed["sizes_pt"]))]
    issues += [C.MSG_COLOR.format(n=n, value=c) for c in sorted(colors - set(allowed["colors"]))]
    issues += [C.MSG_IMAGE.format(n=n, value=p) for h, p in image_hashes(slide.part).items() if h not in allowed["images"]]
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    if not notes.strip():
        issues.append(C.MSG_NOTES_EMPTY.format(n=n))
    return issues


def check_final(final_path, template_path, selections_path=None):
    allowed = inspect_template(template_path)
    slides = list(Presentation(final_path).slides)
    issues = [issue for n, slide in enumerate(slides, start=1) for issue in check_slide(n, slide, allowed)]
    if selections_path:
        expected = len(json.load(open(selections_path))["slides"])
        if expected != len(slides):
            issues.append(C.MSG_SLIDE_COUNT.format(actual=len(slides), expected=expected))
    return issues


def report(issues, warnings=()):
    for line in warnings:
        print(f"WARN {line}", file=sys.stderr)
    for line in issues:
        print(f"FAIL {line}", file=sys.stderr)
    return 1 if issues else 0


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("validate").add_argument("draft")
    sub.add_parser("selections").add_argument("draft")
    sub.add_parser("inspect").add_argument("template")
    check = sub.add_parser("check")
    check.add_argument("final")
    check.add_argument("template")
    check.add_argument("--selections")
    args = parser.parse_args(argv)

    if args.command == "validate":
        slides = load_draft(args.draft)
        errors, warnings = validate_draft(slides)
        if not errors:
            print(f"PASS {args.draft} ({len(slides)} slides)")
        return report(errors, warnings)
    if args.command == "selections":
        result = read_selections(load_draft(args.draft))
        print(json.dumps(result, indent=2))
        return report([], result["issues"])
    if args.command == "inspect":
        print(json.dumps(inspect_template(args.template), indent=2))
        return 0
    issues = check_final(args.final, args.template, args.selections)
    if not issues:
        print("PASS final deck uses only template fonts, sizes, colors, and images; every slide has notes")
    return report(issues)


if __name__ == "__main__":
    sys.exit(main())
