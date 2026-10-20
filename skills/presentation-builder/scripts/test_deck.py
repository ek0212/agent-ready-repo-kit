from pptx import Presentation
from pptx.util import Inches

import deck

VISUALS = "\n".join(f"- Visual {i}. Signals change." for i in range(5))
NOTES = "\n".join(f"- {'word ' * 20}" for _ in range(6)) + "\n- In short: ship small changes often."
DRAFT = f"# Talk\n\n## Slide 1: Why trunk-based\n\n### Visuals\n{VISUALS}\n\n### Speaker notes\n{NOTES}\n"


def test_parse_validate_and_select():
    slides = deck.parse_draft(DRAFT)
    assert slides[0]["heading"] == "Why trunk-based"
    assert deck.validate_draft(slides) == ([], [])
    pruned = deck.parse_draft(DRAFT.replace("- Visual 0. Signals change.\n", ""))
    result = deck.read_selections(pruned)
    assert len(result["slides"][0]["selected_visuals"]) == 4
    assert result["slides"][0]["notes"]


def test_validate_rejects_bad_visuals_and_notes():
    bad = deck.parse_draft(DRAFT.replace(VISUALS, "- One. Two. Three.").replace(NOTES, "Prose notes."))
    errors, _ = deck.validate_draft(bad)
    assert len(errors) == 4


def test_check_flags_non_template_font(tmp_path):
    template, final = tmp_path / "t.pptx", tmp_path / "f.pptx"
    Presentation().save(template)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    run = slide.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text_frame.paragraphs[0].add_run()
    run.text, run.font.name = "x", "Arial"
    prs.save(final)
    assert any("Arial" in issue for issue in deck.check_final(final, template))
